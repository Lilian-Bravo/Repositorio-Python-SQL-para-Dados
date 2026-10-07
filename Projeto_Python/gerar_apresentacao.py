# ==============================================================================
# PROJETO: CONVEX SEGUROS - GERADOR AUTOMATIZADO DA APRESENTAÇÃO EXECUTIVA (PPTX)
# PARCERIA: Blip x Convex Seguros
# PAPEL: Senior Data Strategist
# OBJETIVO: Criar os 10 slides em formato widescreen 16:9, com design moderno,
#           identidade visual da Blip (Deep Navy e Electric Blue) e dados validados.
# ==============================================================================

# ------------------------------------------------------------------------------
# 1. IMPORTAÇÃO DAS BIBLIOTECAS NECESSÁRIAS
# ------------------------------------------------------------------------------
# python-pptx: Biblioteca Python que constrói arquivos do Microsoft PowerPoint (.pptx)
import pptx
from pptx import Presentation

# Medidas de espaçamento e tamanho de fontes:
# - Inches: Mede em polegadas (ex: posição X e Y na tela, largura e altura de caixas)
# - Pt: Mede o tamanho do texto em pontos (ex: tamanho 12, 24, 36)
from pptx.util import Inches, Pt

# Alinhamento de texto (esquerda, centro, direita)
from pptx.enum.text import PP_ALIGN

# Definição de cores personalizadas no formato RGB (Vermelho, Verde, Azul)
from pptx.dml.color import RGBColor

# Formatos geométricos prontos (retângulos, cartões com bordas arredondadas, etc.)
from pptx.enum.shapes import MSO_SHAPE


# ------------------------------------------------------------------------------
# 2. FUNÇÃO PRINCIPAL QUE CRIA A APRESENTAÇÃO
# ------------------------------------------------------------------------------
def criar_apresentacao(nome_arquivo_saida="Apresentacao_Executiva_Convex_Blip.pptx"):
    """
    Esta função monta todos os 10 slides do deck executivo passo a passo.
    Ao final, salva tudo em um arquivo .pptx pronto para abrir no PowerPoint.
    """
    
    # Inicializa um arquivo de PowerPoint em branco
    apresentacao = Presentation()
    
    # Configura o tamanho dos slides para Widescreen (16:9)
    # Largura: 13.333 polegadas | Altura: 7.5 polegadas (Padrão de TVs e projetores modernos)
    apresentacao.slide_width = Inches(13.333)
    apresentacao.slide_height = Inches(7.5)
    
    # O layout número 6 no PowerPoint é o slide 100% em branco (sem caixas de texto prévias)
    layout_em_branco = apresentacao.slide_layouts[6]

    # --------------------------------------------------------------------------
    # DEFINIÇÃO DA PALETA DE CORES OFICIAIS (BLIP & CONVEX SEGUROS)
    # --------------------------------------------------------------------------
    # Cores de Fundo:
    COR_FUNDO_ESCURO = RGBColor(10, 25, 47)      # Deep Navy Blip (#0A192F) - Usado na Capa e Fechamento
    COR_FUNDO_CLARO = RGBColor(248, 250, 252)    # Slate Light (#F8FAFC) - Fundo limpo para leitura
    
    # Cores de Identidade da Blip:
    COR_AZUL_BLIP = RGBColor(0, 102, 255)        # Electric Blue (#0066FF) - Cor primária institucional
    COR_CIANO_BLIP = RGBColor(0, 214, 255)       # Cyber Cyan (#00D6FF) - Destaque em fundos escuros
    
    # Cores de Textos:
    COR_TEXTO_ESCURO = RGBColor(15, 23, 42)      # Quase preto / Slate 900 - Para títulos e leitura
    COR_TEXTO_CINZA = RGBColor(100, 116, 139)    # Cinza médio / Slate 500 - Para explicações secundárias
    COR_TEXTO_BRANCO = RGBColor(255, 255, 255)   # Branco puro
    
    # Cores de Cartões e Bordas:
    COR_FUNDO_CARTAO = RGBColor(255, 255, 255)   # Cartões brancos que parecem flutuar sobre o fundo
    COR_BORDA = RGBColor(226, 232, 240)          # Linha sutil de contorno cinza
    
    # Cores de Destaque Semântico (Indicadores de Negócio):
    COR_VERMELHO_ALERTA = RGBColor(239, 68, 68)  # Destaca o Gargalo de Abandono
    COR_LARANJA_ALERTA = RGBColor(245, 158, 11)  # Destaca o Gargalo de Transbordo Humano
    COR_VERDE_SUCESSO = RGBColor(16, 185, 129)   # Destaca a Resolução 100% pelo Robô (STP)

    # --------------------------------------------------------------------------
    # FUNÇÕES AUXILIARES (FACILITADORES PARA NÃO REPETIR CÓDIGO)
    # --------------------------------------------------------------------------
    def pintar_forma(forma, cor_preenchimento, cor_borda=None, espessura_borda=1):
        """Pinta uma forma geométrica (retângulo, fundo, etc.) sem deixar sobras."""
        forma.fill.solid()
        forma.fill.fore_color.rgb = cor_preenchimento
        if cor_borda:
            forma.line.color.rgb = cor_borda
            forma.line.width = Pt(espessura_borda)
        else:
            forma.line.fill.background()

    def adicionar_cabecalho(slide, titulo, subtitulo_tag="BLIP × CONVEX SEGUROS | BUSINESS CASE DEFENSE", modo_escuro=False):
        """Cria o padrão de cabeçalho com a tag de identificação no topo e o título principal."""
        caixa_texto = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        moldura = caixa_texto.text_frame
        moldura.word_wrap = True
        moldura.margin_left = moldura.margin_top = moldura.margin_right = moldura.margin_bottom = 0
        
        # Tag superior (pequena e colorida)
        p0 = moldura.paragraphs[0]
        p0.text = subtitulo_tag.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = COR_CIANO_BLIP if modo_escuro else COR_AZUL_BLIP
        
        # Título principal do slide
        p1 = moldura.add_paragraph()
        p1.text = titulo
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = COR_TEXTO_BRANCO if modo_escuro else COR_TEXTO_ESCURO

    def criar_cartao(slide, esquerda, topo, largura, altura, cor_fundo=COR_FUNDO_CARTAO, cor_borda=COR_BORDA):
        """Desenha um cartão retangular com cantos levemente arredondados."""
        forma = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, esquerda, topo, largura, altura)
        pintar_forma(forma, cor_fundo, cor_borda)
        return forma

    # ==========================================================================
    # SLIDE 1: CAPA EXECUTIVA (MODO ESCURO DE IMPACTO)
    # ==========================================================================
    slide_1 = apresentacao.slides.add_slide(layout_em_branco)
    
    # 1. Pinta o fundo de Navy Escuro
    fundo_1 = slide_1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    pintar_forma(fundo_1, COR_FUNDO_ESCURO)
    
    # 2. Tag pequena de destaque
    tag_capa = slide_1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(3.6), Inches(0.4))
    pintar_forma(tag_capa, RGBColor(16, 42, 77), COR_AZUL_BLIP)
    p_tag = tag_capa.text_frame.paragraphs[0]
    p_tag.text = "PROCESSO SELETIVO • SENIOR DATA STRATEGIST"
    p_tag.font.size = Pt(9)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COR_CIANO_BLIP
    p_tag.alignment = PP_ALIGN.CENTER

    # 3. Título Principal da Apresentação
    caixa_titulo = slide_1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(2.5))
    moldura_titulo = caixa_titulo.text_frame
    moldura_titulo.word_wrap = True
    
    p_tit = moldura_titulo.paragraphs[0]
    p_tit.text = "Transformação da Jornada Conversacional de Sinistros"
    p_tit.font.size = Pt(36)
    p_tit.font.bold = True
    p_tit.font.color.rgb = COR_TEXTO_BRANCO

    p_sub = moldura_titulo.add_paragraph()
    p_sub.text = "Diagnóstico Analítico do Funil, Governança de Dados e Roadmap de Evolução com IA Generativa"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)
    p_sub.space_before = Pt(14)

    # 4. Três cartões de destaque no rodapé da capa com os números principais
    metricas_capa = [
        ("46,9% STP", "Resolução 100% Digital", COR_AZUL_BLIP),
        ("R$ 157,9k / ano", "Custo em Transbordo Humano", COR_LARANJA_ALERTA),
        ("R$ 57,6M / ano", "Sinistros Abandonados em Risco", COR_VERMELHO_ALERTA)
    ]
    for indice, (valor, rotulo, cor_destaque) in enumerate(metricas_capa):
        pos_x = Inches(0.8 + indice * 3.9)
        cartao_metrica = criar_cartao(slide_1, pos_x, Inches(4.5), Inches(3.6), Inches(1.5), RGBColor(17, 34, 64), cor_destaque)
        moldura_c = cartao_metrica.text_frame
        moldura_c.word_wrap = True
        
        p_val = moldura_c.paragraphs[0]
        p_val.text = valor
        p_val.font.size = Pt(22)
        p_val.font.bold = True
        p_val.font.color.rgb = cor_destaque
        
        p_rot = moldura_c.add_paragraph()
        p_rot.text = rotulo
        p_rot.font.size = Pt(11)
        p_rot.font.color.rgb = RGBColor(203, 213, 225)
        p_rot.space_before = Pt(4)

    # 5. Assinatura do autor e data
    caixa_rodape = slide_1.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5))
    p_rodape = caixa_rodape.text_frame.paragraphs[0]
    p_rodape.text = "Candidato: Senior Data Strategist  |  Empresa Parceira: Blip  |  Cliente: Convex Seguros  |  Período: Jan - Abr 2026"
    p_rodape.font.size = Pt(10)
    p_rodape.font.color.rgb = RGBColor(100, 116, 139)


    # ==========================================================================
    # SLIDE 2: AUDITORIA DE DADOS E GOVERNANÇA DE PRODUÇÃO
    # ==========================================================================
    slide_2 = apresentacao.slides.add_slide(layout_em_branco)
    fundo_2 = slide_2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    pintar_forma(fundo_2, COR_FUNDO_CLARO)
    adicionar_cabecalho(slide_2, "Auditoria e Saneamento da Base de Produção", "GOVERNANÇA & ENGENHARIA DE DADOS")

    # 4 Cartões organizados em grade de 2x2 com os tratamentos aplicados nos dados
    cards_auditoria = [
        ("1. Deduplicação Técnica", "7.502 ➔ 7.000 Registros", 
         "Identificação de 502 registros duplicados de ID_Conversa entre 10 e 24 de março, causados por falha de ingestão de logs. Base limpa consolidada em exatamente 7.000 atendimentos únicos.", COR_LARANJA_ALERTA),
        ("2. Padronização Categórica", "Normalização de Desfechos",
         "Correção de desvios de digitação na coluna Status_Final ('RESOLVIDO_BOT', 'Resolvido_bot', 'ESCALADO HUMANO', 'abandono') para 3 categorias estritas padronizadas.", COR_AZUL_BLIP),
        ("3. Formatação Numérica", "Conversão pt-BR ➔ Float",
         "Substituição de vírgulas por pontos nas colunas de 'Tempo_Resolucao_Min' e 'Valor_Estimado_Sinistro_RS' para viabilizar cálculos matemáticos, e estruturação de datas.", COR_VERDE_SUCESSO),
        ("4. Auditoria de Nulos Válidos", "Detecção do Ponto Cego de CSAT",
         "Constatação de que os 1.658 abandonos possuem nulo legítimo em CSAT. Isso comprovou para a diretoria que a média oficial de 4,02 esconde os clientes frustrados.", COR_VERMELHO_ALERTA)
    ]
    for indice, (titulo_card, subtitulo_card, explicacao, cor_lateral) in enumerate(cards_auditoria):
        coluna = indice % 2
        linha = indice // 2
        pos_x = Inches(0.8 + coluna * 5.9)
        pos_y = Inches(1.7 + linha * 2.5)
        
        criar_cartao(slide_2, pos_x, pos_y, Inches(5.6), Inches(2.2), COR_FUNDO_CARTAO, COR_BORDA)
        
        linha_cor = slide_2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos_x + Inches(0.15), pos_y + Inches(0.25), Inches(0.1), Inches(1.7))
        pintar_forma(linha_cor, cor_lateral)

        caixa_texto_card = slide_2.shapes.add_textbox(pos_x + Inches(0.4), pos_y + Inches(0.2), Inches(5.0), Inches(1.8))
        moldura_c = caixa_texto_card.text_frame
        moldura_c.word_wrap = True
        moldura_c.margin_left = moldura_c.margin_top = moldura_c.margin_right = moldura_c.margin_bottom = 0
        
        p_t = moldura_c.paragraphs[0]
        p_t.text = titulo_card
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = COR_TEXTO_ESCURO
        
        p_s = moldura_c.add_paragraph()
        p_s.text = subtitulo_card
        p_s.font.size = Pt(11)
        p_s.font.bold = True
        p_s.font.color.rgb = cor_lateral
        p_s.space_before = Pt(2)

        p_desc = moldura_c.add_paragraph()
        p_desc.text = explicacao
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = COR_TEXTO_CINZA
        p_desc.space_before = Pt(6)


    # ==========================================================================
    # SLIDE 3: O FUNIL DE SINISTROS & OS DOIS MACRO-GARGALOS
    # ==========================================================================
    slide_3 = apresentacao.slides.add_slide(layout_em_branco)
    fundo_3 = slide_3.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    pintar_forma(fundo_3, COR_FUNDO_CLARO)
    adicionar_cabecalho(slide_3, "Diagnóstico Quantitativo da Jornada (7.000 Conversas)", "ANÁLISE DE CONVERSÃO & FUNIL")

    forma_tabela = slide_3.shapes.add_table(7, 6, Inches(0.8), Inches(1.7), Inches(8.0), Inches(4.8))
    tabela = forma_tabela.table
    tabela.columns[0].width = Inches(2.2)
    tabela.columns[1].width = Inches(1.1)
    tabela.columns[2].width = Inches(1.1)
    tabela.columns[3].width = Inches(1.2)
    tabela.columns[4].width = Inches(1.2)
    tabela.columns[5].width = Inches(1.2)

    titulos_colunas = ["Etapa da Jornada", "Entraram", "Abandonos", "Escalados", "Conclusão", "Evasão %"]
    for col_idx, texto_cabecalho in enumerate(titulos_colunas):
        celula = tabela.cell(0, col_idx)
        celula.fill.solid()
        celula.fill.fore_color.rgb = COR_FUNDO_ESCURO
        p_cab = celula.text_frame.paragraphs[0]
        p_cab.text = texto_cabecalho
        p_cab.font.size = Pt(10)
        p_cab.font.bold = True
        p_cab.font.color.rgb = COR_TEXTO_BRANCO
        p_cab.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT

    dados_etapas = [
        ("1. Início / Saudação", "7.000", "155", "0", "0", "2,2%"),
        ("2. Validação Apólice ⚠️", "6.845", "246", "1.114", "0", "19,9%"),
        ("3. Coleta de Dados", "5.485", "266", "407", "0", "12,3%"),
        ("4. Upload Fotos/Docs 🚨", "4.812", "780", "329", "0", "23,1%"),
        ("5. Análise Automática", "3.703", "211", "207", "0", "11,3%"),
        ("6. Confirmação / Protocolo", "3.285", "0", "0", "3.285", "0,0% (STP)")
    ]
    for linha_idx, valores_linha in enumerate(dados_etapas):
        for col_idx, valor_celula in enumerate(valores_linha):
            celula = tabela.cell(linha_idx + 1, col_idx)
            celula.fill.solid()
            
            if "⚠️" in valores_linha[0] and col_idx == 3:
                celula.fill.fore_color.rgb = RGBColor(254, 243, 199)
            elif "🚨" in valores_linha[0] and col_idx == 2:
                celula.fill.fore_color.rgb = RGBColor(254, 226, 226)
            elif linha_idx % 2 == 1:
                celula.fill.fore_color.rgb = RGBColor(241, 245, 249)
            else:
                celula.fill.fore_color.rgb = COR_FUNDO_CARTAO
            
            p_cel = celula.text_frame.paragraphs[0]
            p_cel.text = valor_celula
            p_cel.font.size = Pt(9.5)
            p_cel.font.bold = True if col_idx in [0, 5] or "⚠️" in valor_celula or "🚨" in valor_celula else False
            p_cel.font.color.rgb = COR_TEXTO_ESCURO
            p_cel.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT

    cartao_lateral_1 = criar_cartao(slide_3, Inches(9.1), Inches(1.7), Inches(3.4), Inches(2.25), COR_FUNDO_CARTAO, COR_LARANJA_ALERTA)
    tf1 = cartao_lateral_1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "⚠️ GARGALO #1: TRANSIÇÃO"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COR_LARANJA_ALERTA
    p1 = tf1.add_paragraph()
    p1.text = "54,2% dos Transbordos"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = COR_TEXTO_ESCURO
    p2 = tf1.add_paragraph()
    p2.text = "1.114 de todos os 2.057 escalonamentos para atendentes humanos acontecem logo na Validação de Apólice, com foco em Vida e Empresas."
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COR_TEXTO_CINZA
    p2.space_before = Pt(4)

    cartao_lateral_2 = criar_cartao(slide_3, Inches(9.1), Inches(4.25), Inches(3.4), Inches(2.25), COR_FUNDO_CARTAO, COR_VERMELHO_ALERTA)
    tf2 = cartao_lateral_2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🚨 GARGALO #2: EVASÃO"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COR_VERMELHO_ALERTA
    p1 = tf2.add_paragraph()
    p1.text = "47,0% dos Abandonos"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = COR_TEXTO_ESCURO
    p2 = tf2.add_paragraph()
    p2.text = "780 dos 1.658 abandonos totais acontecem na etapa de Upload de Fotos e Documentos. O cliente preenche os dados do sinistro, mas desiste no envio de arquivos."
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COR_TEXTO_CINZA
    p2.space_before = Pt(4)


    # ==========================================================================
    # SLIDE 4: DIAGNÓSTICO PROFUNDO DAS CAUSAS-RAIZ
    # ==========================================================================
    slide_4 = apresentacao.slides.add_slide(layout_em_branco)
    fundo_4 = slide_4.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    pintar_forma(fundo_4, COR_FUNDO_CLARO)
    adicionar_cabecalho(slide_4, "Diagnóstico Fino de Causa-Raiz Técnica e Cognitiva", "ANÁLISE DE PRODUTO & COMPORTAMENTO")

    card_cr1 = criar_cartao(slide_4, Inches(0.8), Inches(1.7), Inches(5.7), Inches(4.9), COR_FUNDO_CARTAO, COR_BORDA)
    tf_cr1 = card_cr1.text_frame
    tf_cr1.word_wrap = True
    p = tf_cr1.paragraphs[0]
    p.text = "GARGALO 1: VALIDAÇÃO DE APÓLICE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COR_LARANJA_ALERTA
    
    causas_validacao = [
        ("Rigidez de Titularidade (Seguro Vida):", "No seguro de Vida, o sinistro é acionado pelo beneficiário ou herdeiro. Como o CPF informado não é o do titular falecido, o robô rejeita e transborda para o humano (37,1% de transbordo em Vida)."),
        ("Complexidade Corporativa (Frotas):", "Motoristas de empresas e funcionários acionam apólices jurídicas sem saber os dados do CNPJ tomador, travando o fluxo (apenas 37,5% de sucesso em Corporativo)."),
        ("Intolerância de Entrada de Dados:", "Erros simples de digitação na placa (com ou sem hífen) e pontuação de CPF causam rejeição automática em vez de buscas inteligentes aproximadas.")
    ]
    for titulo_item, descricao_item in causas_validacao:
        p_item = tf_cr1.add_paragraph()
        p_item.text = "• " + titulo_item + " "
        p_item.font.size = Pt(10)
        p_item.font.bold = True
        p_item.font.color.rgb = COR_TEXTO_ESCURO
        p_item.space_before = Pt(10)
        
        texto_corrido = p_item.add_run()
        texto_corrido.text = descricao_item
        texto_corrido.font.bold = False
        texto_corrido.font.color.rgb = COR_TEXTO_CINZA

    card_cr2 = criar_cartao(slide_4, Inches(6.8), Inches(1.7), Inches(5.7), Inches(4.9), COR_FUNDO_CARTAO, COR_BORDA)
    tf_cr2 = card_cr2.text_frame
    tf_cr2.word_wrap = True
    p = tf_cr2.paragraphs[0]
    p.text = "GARGALO 2: UPLOAD DE FOTOS E DOCUMENTOS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COR_VERMELHO_ALERTA
    
    causas_upload = [
        ("Sobrecarga Mental e Falta de Instrução:", "O robô pede vários documentos e fotos de danos de uma só vez, sem dar exemplos visuais de ângulos aceitos nem mostrar uma barra de progresso."),
        ("Arquivos Pesados e Formatos Rejeitados:", "Fotos modernas em alta resolução (como formato HEIC do iPhone ou arquivos acima de 15MB) estouram o limite da API do WhatsApp e dão erro sem explicação ao cliente."),
        ("Expiração de Sessão por Inatividade (Timeout):", "Enquanto o cliente sai do aplicativo para buscar documentos físicos (CNH, documento do carro) ou fotografar o veículo, o tempo limite expira e a conversa é encerrada.")
    ]
    for titulo_item, descricao_item in causas_upload:
        p_item = tf_cr2.add_paragraph()
        p_item.text = "• " + titulo_item + " "
        p_item.font.size = Pt(10)
        p_item.font.bold = True
        p_item.font.color.rgb = COR_TEXTO_ESCURO
        p_item.space_before = Pt(10)
        
        texto_corrido = p_item.add_run()
        texto_corrido.text = descricao_item
        texto_corrido.font.bold = False
        texto_corrido.font.color.rgb = COR_TEXTO_CINZA


    # ==========================================================================
    # SLIDE 5: IMPACTO OPERACIONAL (TMA & OPEX DA CENTRAL)
    # ==========================================================================
    slide_5 = apresentacao.slides.add_slide(layout_em_branco)
    fundo_5 = slide_5.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    pintar_forma(fundo_5, COR_FUNDO_CLARO)
    adicionar_cabecalho(slide_5, "Impacto Operacional & Custo de Transbordo (OPEX)", "MODELAGEM FINANCEIRA & OPERAÇÕES")

    c_tma_bot = criar_cartao(slide_5, Inches(0.8), Inches(1.7), Inches(5.7), Inches(2.2), COR_FUNDO_CARTAO, COR_VERDE_SUCESSO)
    tf_b = c_tma_bot.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "TEMPO MÉDIO DE ATENDIMENTO — VIA BOT"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COR_VERDE_SUCESSO
    p1 = tf_b.add_paragraph()
    p1.text = "9,37 minutos"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = COR_TEXTO_ESCURO
    p2 = tf_b.add_paragraph()
    p2.text = "Resolução 100% digital, rápida e conveniente para o segurado (3.285 casos resolvidos)."
    p2.font.size = Pt(10)
    p2.font.color.rgb = COR_TEXTO_CINZA
    p2.space_before = Pt(4)

    c_tma_humano = criar_cartao(slide_5, Inches(6.8), Inches(1.7), Inches(5.7), Inches(2.2), COR_FUNDO_CARTAO, COR_VERMELHO_ALERTA)
    tf_h = c_tma_humano.text_frame
    tf_h.word_wrap = True
    p = tf_h.paragraphs[0]
    p.text = "TEMPO MÉDIO DE ATENDIMENTO — ESCALADO HUMANO"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COR_VERMELHO_ALERTA
    p1 = tf_h.add_paragraph()
    p1.text = "38,39 minutos (+309%)"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = COR_TEXTO_ESCURO
    p2 = tf_h.add_paragraph()
    p2.text = "O atendente gasta 4x mais tempo porque o cliente tem que repetir tudo o que já havia digitado."
    p2.font.size = Pt(10)
    p2.font.color.rgb = COR_TEXTO_CINZA
    p2.space_before = Pt(4)

    criar_cartao(slide_5, Inches(0.8), Inches(4.15), Inches(11.7), Inches(2.5), COR_FUNDO_CARTAO, COR_BORDA)
    caixa_calculos = slide_5.shapes.add_textbox(Inches(1.0), Inches(4.25), Inches(11.3), Inches(2.2))
    tf_calc = caixa_calculos.text_frame
    tf_calc.word_wrap = True
    
    p = tf_calc.paragraphs[0]
    p.text = "CONSUMO DE CAPACIDADE HUMANA & CUSTO ANUALIZADO (FATOR DE PROJEÇÃO × 3)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COR_AZUL_BLIP

    pontos_financeiros = [
        "• Horas de Analistas Consumidas: 1.316,3 horas nos 4 meses auditados ➔ Projeção de 3.948,8 horas/ano gastas com transbordos.",
        "• Custo Médio por Transbordo: R$ 25,60 gastos em cada atendimento escalado (considerando hora de analista a R$ 40,00 carregada).",
        "• Gasto Operacional Total Atual: R$ 52.650,47 no quadrimestre ➔ R$ 157.951,40 por ano jogados fora por ineficiência do assistente.",
        "• Economia com Melhoria de 40%: Reduzir 823 transbordos por quadrimestre poupa 1.580 horas humanas e economiza R$ 63.180,56/ano em salários."
    ]
    for item in pontos_financeiros:
        p_item = tf_calc.add_paragraph()
        p_item.text = item
        p_item.font.size = Pt(10.5)
        p_item.font.color.rgb = COR_TEXTO_ESCURO
        p_item.space_before = Pt(5)


    # ==========================================================================
    # SLIDE 6: EXPOSIÇÃO EM RISCO & O "PONTO CEGO" DO CSAT
    # ==========================================================================
    slide_6 = apresentacao.slides.add_slide(layout_em_branco)
    fundo_6 = slide_6.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    pintar_forma(fundo_6, COR_FUNDO_CLARO)
    adicionar_cabecalho(slide_6, "Exposição Financeira em Risco & O Ponto Cego do CSAT", "RISCO DE NEGÓCIO & EXPERIÊNCIA DO CLIENTE")

    metricas_risco = [
        ("R$ 19,22M", "Capital Abandonado (4 Meses)", "1.257 sinistros onde o cliente já tinha digitado o valor financeiro do dano e desistiu.", COR_VERMELHO_ALERTA),
        ("R$ 57,65M", "Exposição Anual em Risco", "Volume anual de sinistros em risco de virar reclamação formal no Procon, SUSEP ou processo judicial.", COR_LARANJA_ALERTA),
        ("R$ 44.390", "Ticket Médio Sinistro Vida", "O seguro de Vida retém o maior valor em dinheiro por sinistro, mas sofre a maior taxa de falhas.", COR_AZUL_BLIP)
    ]
    for indice, (valor, titulo_r, explicacao_r, cor_r) in enumerate(metricas_risco):
        c_r = criar_cartao(slide_6, Inches(0.8 + indice * 3.9), Inches(1.7), Inches(3.6), Inches(2.2), COR_FUNDO_CARTAO, cor_r)
        tf_r = c_r.text_frame
        tf_r.word_wrap = True
        
        p = tf_r.paragraphs[0]
        p.text = valor
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = cor_r
        
        p1 = tf_r.add_paragraph()
        p1.text = titulo_r
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COR_TEXTO_ESCURO
        p1.space_before = Pt(4)
        
        p2 = tf_r.add_paragraph()
        p2.text = explicacao_r
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COR_TEXTO_CINZA
        p2.space_before = Pt(4)

    criar_cartao(slide_6, Inches(0.8), Inches(4.15), Inches(11.7), Inches(2.5), COR_FUNDO_CARTAO, COR_BORDA)
    caixa_csat = slide_6.shapes.add_textbox(Inches(1.0), Inches(4.25), Inches(11.3), Inches(2.2))
    tf_csat = caixa_csat.text_frame
    tf_csat.word_wrap = True
    
    p = tf_csat.paragraphs[0]
    p.text = "DESMASCARANDO O 'PONTO CEGO' DO CSAT MÉDIO DE 4,02"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COR_VERMELHO_ALERTA

    itens_csat = [
        ("Queda no CSAT (Bot 4,28 vs. Humano 3,61):", "O cliente escalado para atendimento humano chega irritado com a espera e a repetição de perguntas, derrubando a nota de satisfação em quase 0,7 ponto."),
        ("O Ponto Cego Estatístico:", "A nota de satisfação só é pedida para quem conclui o atendimento. Os 1.658 clientes que abandonaram o chat (23,7% do total) têm ZERO notas registradas."),
        ("Falsa Sensação de Segurança:", "A diretoria achava que o serviço era bem avaliado (nota 4,02), mas a realidade é que quase 1 em cada 4 pessoas desiste frustrada antes do final.")
    ]
    for titulo_c, desc_c in itens_csat:
        p_c = tf_csat.add_paragraph()
        p_c.text = "• " + titulo_c + " "
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = COR_TEXTO_ESCURO
        p_c.space_before = Pt(4)
        
        run_c = p_c.add_run()
        run_c.text = desc_c
        run_c.font.bold = False
        run_c.font.color.rgb = COR_TEXTO_CINZA


    # ==========================================================================
    # SLIDE 7: PRODUTO DE DADOS INTERATIVO COM GENAI (STREAMLIT)
    # ==========================================================================
    slide_7 = apresentacao.slides.add_slide(layout_em_branco)
    fundo_7 = slide_7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    pintar_forma(fundo_7, COR_FUNDO_CLARO)
    adicionar_cabecalho(slide_7, "Cockpit Analítico & Simulador de Negócio com GenAI", "PRODUTO DE DADOS ESCALÁVEL (STREAMLIT + PYTHON)")

    modulos_app = [
        ("Aba 1: Funil & Diagnóstico", "Visualização Dinâmica", 
         "Progressão das 6 etapas com filtros dinâmicos por Canal (WhatsApp/App), Segmento (Varejo/Premium/Corporativo) e Dispositivo (Android/iOS) com gráficos interativos.", COR_AZUL_BLIP),
        ("Aba 2: Simulador Financeiro", "Resiliência à Banca", 
         "Painel interativo com barras deslizantes (sliders) para alterar o Custo/Hora Humana (R$ 25 a R$ 80), Meta de STP e Recuperação de Sinistros em tempo real durante a apresentação.", COR_LARANJA_ALERTA),
        ("Aba 3: Assistente com IA", "Módulo 'Ask the Data'", 
         "Chatbot inteligente acoplado à base que responde perguntas executivas e fornece diagnósticos sintetizados automaticamente usando inteligência artificial generativa.", COR_VERDE_SUCESSO)
    ]
    for indice, (titulo_m, subtitulo_m, descricao_m, cor_m) in enumerate(modulos_app):
        cartao_m = criar_cartao(slide_7, Inches(0.8 + indice * 3.9), Inches(1.7), Inches(3.6), Inches(4.9), COR_FUNDO_CARTAO, cor_m)
        tf_m = cartao_m.text_frame
        tf_m.word_wrap = True
        
        p = tf_m.paragraphs[0]
        p.text = titulo_m
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = cor_m
        
        p1 = tf_m.add_paragraph()
        p1.text = subtitulo_m
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COR_TEXTO_ESCURO
        p1.space_before = Pt(4)
        
        p2 = tf_m.add_paragraph()
        p2.text = descricao_m
        p2.font.size = Pt(10)
        p2.font.color.rgb = COR_TEXTO_CINZA
        p2.space_before = Pt(10)
        
        p3 = tf_m.add_paragraph()
        p3.text = "Status: 100% Funcional e operando na máquina via Streamlit."
        p3.font.size = Pt(9.5)
        p3.font.italic = True
        p3.font.color.rgb = COR_AZUL_BLIP
        p3.space_before = Pt(14)


    # ==========================================================================
    # SLIDE 8: ROADMAP ESTRATÉGICO DE PRODUTO
    # ==========================================================================
    slide_8 = apresentacao.slides.add_slide(layout_em_branco)
    fundo_8 = slide_8.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    pintar_forma(fundo_8, COR_FUNDO_CLARO)
    adicionar_cabecalho(slide_8, "Roadmap Estratégico de Evolução de Produto (2026)", "TRANSFORMAÇÃO DA EXPERIÊNCIA")

    fases_roadmap = [
        ("FASE 1: 0 A 3 MESES", "Quick Wins & Eficiência Imediata", [
            "Guia Visual de Fotos: Mensagens com exemplos visuais no WhatsApp mostrando fotos nítidas aceitas.",
            "Compressão no Celular: Redução automática de tamanho de fotos (HEIC/JPG) no próprio smartphone antes do envio.",
            "Warm Handoff por IA: Resumo automático do sinistro gerado por LLM na tela do atendente humano.",
            "Impacto: Redução imediata do tempo de atendimento de 38 min para <25 min e estabilização de CSAT."
        ], COR_AZUL_BLIP),
        ("FASE 2: 3 A 6 MESES", "Inteligência & Segmentação", [
            "Validação Visual por IA: Visão computacional checando em tempo real se a foto do dano está nítida.",
            "Fluxos para Terceiros: Validação inteligente para beneficiários de Vida e condutores de frotas da empresa.",
            "Régua Ativa no WhatsApp: Mensagem automática após 15 minutos chamando o cliente de volta para terminar o upload.",
            "Impacto: +12 pontos percentuais de STP e corte de 40% nos abandonos de upload."
        ], COR_LARANJA_ALERTA),
        ("FASE 3: 6 A 12 MESES", "Automação Total (End-to-End)", [
            "Fast-Track STP: Aprovação e liberação imediata do pagamento para sinistros pequenos e clientes sem histórico de risco.",
            "Agendamento de Oficinas: O cliente escolhe a oficina credenciada e agenda o reparo diretamente no chat.",
            "Antifraude em Tempo Real: IA analisando histórico e fotos para barrar fraudes antes da aprovação.",
            "Impacto: Resolução digital acima de 65%, economia anual de OPEX > R$ 80k e CSAT superior a 4,6."
        ], COR_VERDE_SUCESSO)
    ]
    for indice, (titulo_f, subtitulo_f, lista_itens, cor_f) in enumerate(fases_roadmap):
        cartao_f = criar_cartao(slide_8, Inches(0.8 + indice * 3.9), Inches(1.7), Inches(3.6), Inches(4.9), COR_FUNDO_CARTAO, cor_f)
        tf_f = cartao_f.text_frame
        tf_f.word_wrap = True
        
        p = tf_f.paragraphs[0]
        p.text = titulo_f
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = cor_f
        
        p1 = tf_f.add_paragraph()
        p1.text = subtitulo_f
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = COR_TEXTO_ESCURO
        p1.space_before = Pt(4)

        for item_f in lista_itens:
            p_item_f = tf_f.add_paragraph()
            p_item_f.text = "• " + item_f
            p_item_f.font.size = Pt(9.5)
            p_item_f.font.color.rgb = COR_TEXTO_CINZA
            p_item_f.space_before = Pt(6)


    # ==========================================================================
    # SLIDE 9: ORQUESTRAÇÃO DE IA & GOVERNANÇA METODOLÓGICA
    # ==========================================================================
    slide_9 = apresentacao.slides.add_slide(layout_em_branco)
    fundo_9 = slide_9.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    pintar_forma(fundo_9, COR_FUNDO_CLARO)
    adicionar_cabecalho(slide_9, "Como a IA Generativa foi Orquestrada no Projeto", "GOVERNANÇA METODOLÓGICA DE IA")

    pilares_ia = [
        ("Ferramentas Selecionadas", "Foundation Models & Python", 
         "Uso de Modelos de Linguagem Avançados (Claude 3.5 Sonnet / GPT-4o / Gemini Pro) integrados com execução Python para engenharia de dados automatizada e síntese analítica.", COR_AZUL_BLIP),
        ("Engenharia de Prompts", "Role-Prompting & Raciocínio Passo a Passo", 
         "Prompts definindo o papel de especialista ('Atue como Senior Data Strategist focado em seguradoras e Blip'), exigindo raciocínio prévio em etapas antes de emitir recomendações.", COR_LARANJA_ALERTA),
        ("Validação Humana (Human-in-the-Loop)", "Detecção de Falhas e Alucinações", 
         "Supervisão humana para auditar as 502 linhas duplicadas no banco de produção, identificar a falta de CSAT nos abandonos e calibrar a projeção anual multiplicando por 3.", COR_VERDE_SUCESSO),
        ("Geração do Produto de Dados", "Automação de Código e Dashboard", 
         "Aceleração do desenvolvimento do painel Streamlit interativo com gráficos Plotly e módulo consultivo em linguagem natural para apoiar a diretoria da Convex.", RGBColor(147, 51, 234))
    ]
    for indice, (titulo_ia, subtitulo_ia, texto_ia, cor_ia) in enumerate(pilares_ia):
        coluna = indice % 2
        linha = indice // 2
        pos_x = Inches(0.8 + coluna * 5.9)
        pos_y = Inches(1.7 + linha * 2.5)
        
        criar_cartao(slide_9, pos_x, pos_y, Inches(5.6), Inches(2.2), COR_FUNDO_CARTAO, COR_BORDA)
        
        barra_cor = slide_9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos_x + Inches(0.15), pos_y + Inches(0.25), Inches(0.1), Inches(1.7))
        pintar_forma(barra_cor, cor_ia)

        caixa_ia = slide_9.shapes.add_textbox(pos_x + Inches(0.4), pos_y + Inches(0.2), Inches(5.0), Inches(1.8))
        tf_ia = caixa_ia.text_frame
        tf_ia.word_wrap = True
        
        p_t = tf_ia.paragraphs[0]
        p_t.text = titulo_ia
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = COR_TEXTO_ESCURO
        
        p_s = tf_ia.add_paragraph()
        p_s.text = subtitulo_ia
        p_s.font.size = Pt(11)
        p_s.font.bold = True
        p_s.font.color.rgb = cor_ia
        p_s.space_before = Pt(2)

        p_desc = tf_ia.add_paragraph()
        p_desc.text = texto_ia
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = COR_TEXTO_CINZA
        p_desc.space_before = Pt(6)


    # ==========================================================================
    # SLIDE 10: CONCLUSÃO & DEFESA DE PREMISSAS PERANTE A BANCA
    # ==========================================================================
    slide_10 = apresentacao.slides.add_slide(layout_em_branco)
    fundo_10 = slide_10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    pintar_forma(fundo_10, COR_FUNDO_ESCURO)
    adicionar_cabecalho(slide_10, "Conclusão Executiva & Prontidão para a Banca", "DEFESA ESTRATÉGICA • 20 MIN + 10 MIN Q&A", modo_escuro=True)

    testes_estresse = [
        ("Mudança ao Vivo: O orçamento para IA foi cortado em 50%. O que manter?",
         "Defesa Estratégica: Priorizamos integralmente as melhorias da Fase 1 (Quick Wins). Guias de fotos no WhatsApp, compressão de imagens no celular e regras mais flexíveis de validação custam quase zero em servidores e resolvem mais de 50% dos abandonos de upload."),
        ("Mudança ao Vivo: A diretoria quer priorizar apenas Seguro Auto e cancelar o resto.",
         "Defesa Estratégica: Ajustamos o foco para a validação visual de danos de lataria e para-choques por visão computacional. O Seguro Auto concentra 54,8% de todos os sinistros (3.835 casos), entregando escala rápida e retorno financeiro imediato."),
        ("Mudança ao Vivo: O custo da hora humana na verdade é R$ 60,00 e não R$ 40,00.",
         "Defesa Estratégica: Demonstramos ao vivo no Simulador que cada transbordo passa a custar R$ 38,40 e o prejuízo anual sobe para R$ 236,9 mil. Isso aumenta o retorno sobre investimento (ROI) e torna a automação ainda mais urgente para a empresa.")
    ]
    for indice, (pergunta_stress, resposta_defesa) in enumerate(testes_estresse):
        c_stress = criar_cartao(slide_10, Inches(0.8), Inches(1.7 + indice * 1.6), Inches(11.7), Inches(1.4), RGBColor(17, 34, 64), COR_AZUL_BLIP)
        tf_stress = c_stress.text_frame
        tf_stress.word_wrap = True
        
        p = tf_stress.paragraphs[0]
        p.text = pergunta_stress
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COR_CIANO_BLIP
        
        p1 = tf_stress.add_paragraph()
        p1.text = resposta_defesa
        p1.font.size = Pt(10)
        p1.font.color.rgb = RGBColor(203, 213, 225)
        p1.space_before = Pt(4)

    caixa_fechamento = slide_10.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.6))
    p_fechamento = caixa_fechamento.text_frame.paragraphs[0]
    p_fechamento.text = "Obrigado! Aberto para perguntas da banca, demonstração do Dashboard e simulação de cenários dinâmicos."
    p_fechamento.font.size = Pt(12)
    p_fechamento.font.bold = True
    p_fechamento.font.color.rgb = COR_TEXTO_BRANCO
    p_fechamento.alignment = PP_ALIGN.CENTER

    apresentacao.save(nome_arquivo_saida)
    print(f"✅ Arquivo salvo com sucesso em: {nome_arquivo_saida}")


# ------------------------------------------------------------------------------
# 3. PONTO DE ENTRADA DO PROGRAMA
# ------------------------------------------------------------------------------
if __name__ == "__main__":
    criar_apresentacao()